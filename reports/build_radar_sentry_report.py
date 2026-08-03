"""Build the editable RMUC radar-sentry decision-system status deck."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "reports" / "assets"
OUT = ROOT / "reports" / "RMUC_Radar_Sentry_Decision_Report_20260731.pptx"

W, H = Inches(13.333), Inches(7.5)
FONT = "Noto Sans CJK SC"
BG = RGBColor(12, 17, 25)
PANEL = RGBColor(23, 31, 43)
PANEL_2 = RGBColor(30, 40, 54)
TEXT = RGBColor(240, 244, 248)
MUTED = RGBColor(166, 180, 195)
RED = RGBColor(225, 77, 77)
BLUE = RGBColor(66, 145, 244)
CYAN = RGBColor(50, 203, 217)
GOLD = RGBColor(245, 187, 71)
GREEN = RGBColor(96, 205, 142)


def set_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def text_box(slide, x, y, w, h, text, *, size=16, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.07):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, color, radius=False, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = line if line else color
    return shape


def title(slide, kicker: str, heading: str, page: int) -> None:
    text_box(slide, Inches(0.55), Inches(0.28), Inches(9.7), Inches(0.22), kicker.upper(), size=9, color=CYAN, bold=True)
    text_box(slide, Inches(0.55), Inches(0.53), Inches(11.9), Inches(0.55), heading, size=25, bold=True)
    rect(slide, Inches(0.55), Inches(1.12), Inches(12.25), Inches(0.02), PANEL_2)
    text_box(slide, Inches(12.25), Inches(0.34), Inches(0.5), Inches(0.25), f"{page:02d}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, heading: str, body: str, accent=CYAN, body_size=13):
    rect(slide, x, y, w, h, PANEL, radius=True)
    rect(slide, x, y, Inches(0.05), h, accent, radius=True)
    text_box(slide, x + Inches(0.18), y + Inches(0.14), w - Inches(0.32), Inches(0.3), heading, size=14, color=accent, bold=True)
    text_box(slide, x + Inches(0.18), y + Inches(0.52), w - Inches(0.32), h - Inches(0.64), body, size=body_size, color=TEXT)


def bullet_lines(slide, x, y, w, lines, *, size=14, color=TEXT, bullet_color=CYAN, gap=0.38):
    for i, line in enumerate(lines):
        yy = y + Inches(i * gap)
        text_box(slide, x, yy, Inches(0.2), Inches(0.24), "•", size=size + 3, color=bullet_color, bold=True)
        text_box(slide, x + Inches(0.23), yy + Inches(0.02), w - Inches(0.23), Inches(0.3), line, size=size, color=color)


def add_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    map_path = ASSETS / "semantic_map_aligned.png"
    slide.shapes.add_picture(str(map_path), Inches(7.25), Inches(0), Inches(6.08), Inches(7.5))
    rect(slide, Inches(6.55), Inches(0), Inches(6.8), Inches(7.5), BG)
    rect(slide, Inches(0.62), Inches(0.75), Inches(0.09), Inches(1.4), RED, radius=True)
    text_box(slide, Inches(0.93), Inches(0.78), Inches(5.35), Inches(0.26), "RMUC 2026  |  哨兵整车系统", size=12, color=CYAN, bold=True)
    text_box(slide, Inches(0.93), Inches(1.16), Inches(5.75), Inches(1.35), "哨兵整车系统\n决策端搭建进展汇报", size=30, bold=True)
    text_box(slide, Inches(0.95), Inches(2.7), Inches(5.25), Inches(0.6), "感知 / 导航 / 自瞄 / 裁判通信 · 语义地图 · 在线 PPO", size=15, color=MUTED)
    card(slide, Inches(0.95), Inches(4.15), Inches(5.45), Inches(1.35), "本次汇报重点", "完整展示整车执行链路；重点汇报雷达站高层战术决策、语义地图、规则世界模型和在线 PPO 的搭建进度。", CYAN, 14)
    text_box(slide, Inches(0.95), Inches(6.62), Inches(5.4), Inches(0.25), "汇报人：俞典  ·  2026-07-31  ·  内部技术汇报", size=10, color=MUTED)


def add_scope(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_background(slide); title(slide, "01 / Whole Vehicle", "整车系统分层：决策端是上层，不取代车端", 2)
    card(slide, Inches(0.65), Inches(1.55), Inches(3.8), Inches(3.8), "雷达站战术决策（重点）", "• 汇总全场结构化状态\n• 选择 goal / target / fire mode\n• 评估代价地图、局势与规则\n• 触发低血量回补给等安全 skill", RED, 15)
    card(slide, Inches(4.78), Inches(1.55), Inches(3.8), Inches(3.8), "导航与底盘执行", "• 全局 JPS / 局部 ESDF / MINCO\n• 局部控制与动态避障\n• cmd_vel 仲裁与执行安全\n• 不可达、碰撞、热量和弹药保护", BLUE, 15)
    card(slide, Inches(8.91), Inches(1.55), Inches(3.75), Inches(3.8), "感知、自瞄与通信", "• FAST-LIO / ROG-Map / 语义地图\n• tracker / 火控 / 云台执行\n• 裁判系统与 CDC 状态协议\n• 雷达失联后的本地兜底", GOLD, 15)
    text_box(slide, Inches(0.72), Inches(5.95), Inches(11.8), Inches(0.5), "学习模块只输出战术意图；局部规划、自瞄与硬安全始终在车端闭环。", size=18, color=GREEN, bold=True, align=PP_ALIGN.CENTER)


def add_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_background(slide); title(slide, "02 / Architecture", "雷达站端到端决策链路", 3)
    nodes = [
        ("雷达 / 裁判态势", "敌我坐标、HP、时间、建筑、金币", CYAN),
        ("决策与语义规则", "候选目标、PPO、recovery、安全 skill", RED),
        ("Radar Tactical State", "99 B 固定帧 / TTL / 优先级 / fallback", GREEN),
        ("车端仲裁器", "一级 action / 二级态势 / 三级本地兜底", GOLD),
        ("导航与底盘", "JPS、MINCO、ESDF、cmd_vel mux", BLUE),
        ("感知与自瞄", "定位建图、tracker、火控、云台", CYAN),
    ]
    x0, y = Inches(0.6), Inches(2.2)
    for i, (head, body, col) in enumerate(nodes):
        x = x0 + Inches(i * 2.08)
        rect(slide, x, y, Inches(1.72), Inches(1.45), PANEL, radius=True)
        text_box(slide, x + Inches(0.1), y + Inches(0.18), Inches(1.52), Inches(0.35), head, size=13, color=col, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, x + Inches(0.1), y + Inches(0.65), Inches(1.52), Inches(0.6), body, size=10, color=TEXT, align=PP_ALIGN.CENTER)
        if i < len(nodes) - 1:
            text_box(slide, x + Inches(1.75), y + Inches(0.55), Inches(0.28), Inches(0.3), "→", size=20, color=MUTED, align=PP_ALIGN.CENTER)
    text_box(slide, Inches(1.2), Inches(4.55), Inches(10.9), Inches(0.5), "闭环：雷达站下发战术意图；车端仲裁、导航与自瞄反馈路径、伤害、死亡和执行失败原因。", size=16, color=MUTED, align=PP_ALIGN.CENTER)


def add_vehicle_stack(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_background(slide); title(slide, "03 / Vehicle Stack", "Gazebo 整车执行底座：已有模块与接口", 4)
    cards = [
        ("定位与地图", "FAST-LIO TF 链；ROG-Map、/map、/local_grid_map、/local_esdf。", CYAN),
        ("全局与局部导航", "/navigation/goal_pose → JPS global path → MINCO trajectory → ESDF local controller。", BLUE),
        ("底盘仲裁", "cmd_vel_mux 在 teleop / navigation / timeout_stop 之间仲裁；车端唯一发布 /cmd_vel。", GREEN),
        ("直接火控仿真", "TrackedUnits → FireControlNode → serial/send → autoaim_bridge → Gazebo 云台关节。", RED),
        ("三层决策仲裁", "Radar Action / Radar State / Local Only；TTL、滞回、冷却和本地硬安全。", GOLD),
        ("本项目接入点", "向仲裁器提供 TacticalAction 与雷达态势；不直接越过 ROS 导航与自瞄节点。", PANEL_2),
    ]
    for i, (h, b, c) in enumerate(cards):
        x = Inches(0.7 + (i % 3) * 4.17); y = Inches(1.5 + (i // 3) * 2.02)
        card(slide, x, y, Inches(3.78), Inches(1.53), h, b, c, 12)


def add_map(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_background(slide); title(slide, "04 / Semantic Map", "黑白可通行图与语义层分离", 5)
    slide.shapes.add_picture(str(ASSETS / "semantic_map_aligned.png"), Inches(0.62), Inches(1.45), Inches(7.2), Inches(4.0))
    card(slide, Inches(8.15), Inches(1.45), Inches(4.45), Inches(1.1), "硬约束层", "黑白障碍、车体可通行性、A* 禁止穿越。", RED, 13)
    card(slide, Inches(8.15), Inches(2.75), Inches(4.45), Inches(1.1), "语义与收益层", "补给、基地、前哨、堡垒、中央高地、隧道等独立 JSON 多边形。", CYAN, 13)
    card(slide, Inches(8.15), Inches(4.05), Inches(4.45), Inches(1.1), "标定与使用", "地图坐标为 28 × 15 m；语义区服务于规则、reward、cost 与 Foxglove 可视化。", GOLD, 13)
    text_box(slide, Inches(0.72), Inches(5.85), Inches(11.9), Inches(0.35), "原则：语义区不覆盖硬障碍；地图收益必须通过正式规则触发，不把颜色当作奖励。", size=15, color=GREEN, bold=True, align=PP_ALIGN.CENTER)


def add_candidates(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_background(slide); title(slide, "05 / Target Selection", "车辆目标选择：从“分类输出”升级为候选成本评估", 6)
    left = ["目标是否存活、是否可攻击", "A* 到合法射击站位的可达性", "真实路径代价与动态威胁风险", "射击站位距离与预期有效 DPS", "目标 HP、武器威胁、基地攻击压力", "己方/敌方前哨、时间、低血量 recovery 状态"]
    bullet_lines(slide, Inches(0.75), Inches(1.55), Inches(5.3), left, size=14, gap=0.53)
    rect(slide, Inches(6.35), Inches(1.45), Inches(5.9), Inches(3.75), PANEL, radius=True)
    text_box(slide, Inches(6.65), Inches(1.75), Inches(5.25), Inches(0.35), "每辆敌方地面车的候选特征", size=18, color=CYAN, bold=True)
    rows = [("可达", "A* standoff path"), ("风险", "path risk / enemy threat"), ("收益", "expected DPS / target HP"), ("紧急度", "base / outpost pressure"), ("约束", "视线、射程、弹药、热量")]
    for i, (a, b) in enumerate(rows):
        yy = Inches(2.28 + i * 0.52)
        rect(slide, Inches(6.65), yy, Inches(1.2), Inches(0.33), PANEL_2, radius=True)
        text_box(slide, Inches(6.72), yy + Inches(0.03), Inches(1.05), Inches(0.22), a, size=11, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, Inches(8.05), yy + Inches(0.03), Inches(3.55), Inches(0.24), b, size=12, color=TEXT)
    text_box(slide, Inches(0.75), Inches(5.72), Inches(11.7), Inches(0.45), "状态：候选特征与互斥中央高地规则已实现；新网络需全新训练，当前已暂停以优先完成汇报。", size=14, color=MUTED, align=PP_ALIGN.CENTER)


def add_world(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_background(slide); title(slide, "06 / World Model", "1 Hz 规则世界模型：训练高层战术，不替代实车物理", 7)
    cards = [
        ("交战", "车辆有效 DPS：3 m / 5 m 分层；英雄保留 4 秒 42 mm 节奏；地面单位不能攻击空中。", RED),
        ("建筑", "前哨存活时基地无敌；建筑伤害要求停留 2 秒且 5 m 内无敌方地面车。", GOLD),
        ("复活", "读条与比赛已进行时间相关；补给区或基地低血量时读条 4 倍。", CYAN),
        ("回血", "补给区 10% HP/s；4 分钟后脱战 6 秒可达 25% HP/s。", GREEN),
        ("陪练", "11 个冻结 offlineRL 角色给出路线/建筑意图；贴身合法交战由 reactive auto-aim 补齐。", BLUE),
        ("边界", "金币、远程回血、立即复活、虚弱解除与细粒度弹道尚未完整接入。", PANEL_2),
    ]
    for i, (h, b, c) in enumerate(cards):
        x = Inches(0.7 + (i % 3) * 4.17); y = Inches(1.55 + (i // 3) * 2.05)
        card(slide, x, y, Inches(3.78), Inches(1.55), h, b, c, 12)


def add_skills(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_background(slide); title(slide, "07 / Safety Skills", "低血量回补给：规则驱动的安全 skill", 8)
    text_box(slide, Inches(0.75), Inches(1.48), Inches(3.2), Inches(0.28), "HP ≤ 25%", size=17, color=RED, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, Inches(4.75), Inches(1.48), Inches(3.2), Inches(0.28), "锁定 healing_2", size=17, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, Inches(8.75), Inches(1.48), Inches(3.2), Inches(0.28), "HP ≥ 90%", size=17, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    for x, col in ((Inches(1.35), RED), (Inches(5.35), CYAN), (Inches(9.35), GREEN)):
        rect(slide, x, Inches(2.05), Inches(2.0), Inches(1.15), PANEL, radius=True)
        rect(slide, x + Inches(0.73), Inches(2.38), Inches(0.55), Inches(0.55), col, radius=True)
    text_box(slide, Inches(3.48), Inches(2.35), Inches(1.2), Inches(0.35), "→", size=28, color=MUTED, align=PP_ALIGN.CENTER)
    text_box(slide, Inches(7.48), Inches(2.35), Inches(1.2), Inches(0.35), "→", size=28, color=MUTED, align=PP_ALIGN.CENTER)
    card(slide, Inches(0.95), Inches(4.15), Inches(11.45), Inches(1.2), "为什么这是安全 skill 而不是让 PPO 自己猜", "旧策略曾在 80 HP 停留中央高地。现在 recovery 在观测 mask 与执行层同时生效；PPO 不会因为动作被静默改写而学到错误因果。", GOLD, 14)


def add_gazebo_contract(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_background(slide); title(slide, "08 / Decision Interface", "与 Gazebo 三层决策蓝图的接口：雷达站负责一级", 9)
    card(slide, Inches(0.7), Inches(1.45), Inches(3.85), Inches(2.15), "一级：Radar Action", "本项目输出 TacticalAction：GO_TO / PURSUE / DEFEND / RETREAT / SUPPORT，带 target、goal、约束、TTL、优先级和 fallback。", RED, 14)
    card(slide, Inches(4.75), Inches(1.45), Inches(3.85), Inches(2.15), "二级：Radar State", "action 过期但态势仍新鲜时，哨兵端行为树 + 效用函数读取敌我位置、HP 与语义地图自主排序。", CYAN, 14)
    card(slide, Inches(8.8), Inches(1.45), Inches(3.85), Inches(2.15), "三级：Local Only", "雷达站超时后，Gazebo/车端依赖视觉、激光、里程计、本地 ESDF 与语义地图保守兜底。", GOLD, 14)
    rect(slide, Inches(0.95), Inches(4.15), Inches(11.55), Inches(1.22), PANEL, radius=True)
    text_box(slide, Inches(1.2), Inches(4.37), Inches(10.9), Inches(0.3), "RadarTacticalStateWireV1（99 B，CRC、序列号、TTL）", size=18, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, Inches(1.2), Inches(4.78), Inches(10.9), Inches(0.28), "CDC / 裁判交互链路 → latest_radar_state → 车端仲裁器 → /navigation/goal_pose 与自瞄优先级", size=13, color=TEXT, align=PP_ALIGN.CENTER)
    text_box(slide, Inches(0.78), Inches(5.9), Inches(11.8), Inches(0.3), "车端已有导航：/navigation/goal_pose → JPS 全局路径 → MINCO/ESDF 局部控制 → /cmd_vel；雷达站不直接发布底盘速度。", size=14, color=MUTED, align=PP_ALIGN.CENTER)


def add_replay(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_background(slide); title(slide, "09 / Replay Evidence", "统一标注回放：用于验证链路，不用于宣称真实胜率", 10)
    slide.shapes.add_picture(str(ASSETS / "replay_rules_frame.png"), Inches(0.62), Inches(1.4), Inches(7.35), Inches(4.08))
    card(slide, Inches(8.28), Inches(1.42), Inches(4.25), Inches(1.05), "统一渲染合同", "无网格比赛底图；固定字体、血条、建筑 HP、射程环、意图线和伤害线。", CYAN, 12)
    card(slide, Inches(8.28), Inches(2.7), Inches(4.25), Inches(1.05), "可核验数据", "steps.csv、units.csv、combat_events.csv、summary.json 与 MP4 同步生成。", GOLD, 12)
    card(slide, Inches(8.28), Inches(3.98), Inches(4.25), Inches(1.05), "本次规则回放", "420 秒；蓝方会贴身反击；红哨兵触发多段回补给。旧 PPO 仅作规则回归，不作最终策略。", RED, 12)
    text_box(slide, Inches(0.72), Inches(5.86), Inches(11.8), Inches(0.35), "回放目的：发现“贴脸不打、早期推塔、低血量不回补给”等系统错误，再反馈给世界模型和训练。", size=14, color=GREEN, bold=True, align=PP_ALIGN.CENTER)


def add_status(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_background(slide); title(slide, "10 / Status", "当前完成项、已发现问题与结论边界", 11)
    card(slide, Inches(0.68), Inches(1.45), Inches(3.9), Inches(3.9), "已完成", "语义地图对齐与可视化\n0.1 m A* 与成本图\n11 角色 offlineRL 陪练接入\n规则交战、建筑、复活、补给\n主动猎杀 reward / 低血量 recovery\n标准回放与 telemetry", GREEN, 14)
    card(slide, Inches(4.72), Inches(1.45), Inches(3.9), Inches(3.9), "已暴露并修正", "陪练 target=none 导致贴脸不打\n建筑目标曾绕过 2 秒停留规则\n复活后错误清空弹药\n低血量错误离开补给区\n中央高地曾被建模为无条件收益", GOLD, 14)
    card(slide, Inches(8.76), Inches(1.45), Inches(3.9), Inches(3.9), "不能过度解读", "单局 red_win / blue_win\n离线 held-out 分数\n世界模型中的单一陪练池\n未标定的实车速度、命中率、通信误差\n尚未训练完成的候选目标版本", RED, 14)


def add_roadmap(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_background(slide); title(slide, "11 / Roadmap", "下一阶段：先稳定对手，再训练主哨兵，再接入 Gazebo/实车", 12)
    steps = [
        ("A", "冻结规则", "补齐金币、远程回血、立即复活、弱化与高地占领细节；固定回归场景。", GOLD),
        ("B", "稳定陪练池", "按角色与风格做 leave-one-style-out；检查路线、交战、建筑推进分布。", BLUE),
        ("C", "候选目标 PPO", "在主动反击陪练环境重新训练；比较基地守护、有效 DPS、胜率与熵。", RED),
        ("D", "实车标定", "接入雷达状态、车端导航、射击命中和通信延迟；收敛仿真参数。", GREEN),
    ]
    for i, (tag, h, b, c) in enumerate(steps):
        y = Inches(1.48 + i * 1.16)
        rect(slide, Inches(0.8), y, Inches(0.55), Inches(0.55), c, radius=True)
        text_box(slide, Inches(0.8), y + Inches(0.12), Inches(0.55), Inches(0.2), tag, size=13, color=BG, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, Inches(1.55), y + Inches(0.03), Inches(2.2), Inches(0.28), h, size=17, color=c, bold=True)
        text_box(slide, Inches(3.72), y + Inches(0.03), Inches(8.3), Inches(0.42), b, size=14, color=TEXT)


def add_close(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_background(slide)
    rect(slide, Inches(0.65), Inches(0.8), Inches(0.08), Inches(2.0), CYAN, radius=True)
    text_box(slide, Inches(0.98), Inches(0.92), Inches(10.5), Inches(0.3), "核心结论", size=13, color=CYAN, bold=True)
    text_box(slide, Inches(0.98), Inches(1.35), Inches(10.6), Inches(1.1), "雷达站端决策系统已经具备\n“可运行、可回放、可校验”的最小闭环。", size=29, bold=True)
    text_box(slide, Inches(1.0), Inches(3.15), Inches(10.6), Inches(0.7), "下一步不再只是增加 reward，而是用稳定规则、真实陪练分布和逐目标路径代价，让哨兵学会在进攻、回补给与基地防守之间做长期选择。", size=17, color=MUTED)
    rect(slide, Inches(0.98), Inches(5.45), Inches(11.25), Inches(0.75), PANEL, radius=True)
    text_box(slide, Inches(1.15), Inches(5.68), Inches(10.9), Inches(0.25), "雷达站高层决策  ×  车端传统执行  ×  可验证规则世界模型", size=16, color=GREEN, bold=True, align=PP_ALIGN.CENTER)


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    add_cover(prs); add_scope(prs); add_architecture(prs); add_vehicle_stack(prs); add_map(prs)
    add_candidates(prs); add_world(prs); add_skills(prs); add_gazebo_contract(prs); add_replay(prs)
    add_status(prs); add_roadmap(prs); add_close(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
